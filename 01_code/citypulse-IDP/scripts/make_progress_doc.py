"""Creates/updates Week1_Progress.docx — a concise, append-style progress log.
Not part of the research deliverables; a working tracker only. Kept out of git (.gitignore).
Run: python scripts/make_progress_doc.py init   -> creates the file with the spike skeleton
     python scripts/make_progress_doc.py log "<task>" "<one-line update>" [status]
"""
import sys
import datetime
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOC_PATH = Path(__file__).resolve().parent.parent / "Week1_Progress.docx"

SPIKES = [
    ("T0.1", "SLM on cheap Android hardware", "BLOCKED — needs a physical <=15k Android device; cannot be run by an agent."),
    ("T0.2", "Graph build + query time", "IN PROGRESS"),
    ("T0.3", "Ethics/IEC application submitted", "BLOCKED — needs a faculty PI to file; cannot be done by an agent."),
    ("T0.4", "Accounts + Tier-2 data requests claimed", "BLOCKED — needs a human to create accounts / sign applications."),
    ("T0.5", "Watchlist feasibility spike", "IN PROGRESS"),
]

def init_doc():
    doc = Document()
    title = doc.add_heading("CityPulse IDP — Week 1 Progress", level=1)
    sub = doc.add_paragraph(f"Started {datetime.date.today().isoformat()}. Concise, append-only log — newest entries at the bottom of each section.")
    sub.runs[0].italic = True

    doc.add_heading("Spike status (at a glance)", level=2)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Task", "What", "Status"
    for tid, what, status in SPIKES:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = tid, what, status

    doc.add_heading("Log", level=2)
    doc.add_paragraph("(entries added as work happens)")
    doc.save(DOC_PATH)
    print(f"Created {DOC_PATH}")

def log_entry(task: str, update: str, status: str | None = None):
    doc = Document(DOC_PATH)

    # Update the status table if a status was given
    if status:
        for table in doc.tables:
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == task:
                    row.cells[2].text = status

    # Find the "Log" heading and append after the last paragraph in that section
    log_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Log":
            log_idx = i
            break

    ts = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
    p = doc.add_paragraph()
    run = p.add_run(f"[{ts}] {task} — ")
    run.bold = True
    p.add_run(update)

    doc.save(DOC_PATH)
    print(f"Logged: [{ts}] {task} — {update}")

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    cmd = sys.argv[1]
    if cmd == "init":
        init_doc()
    elif cmd == "log":
        task = sys.argv[2]
        update = sys.argv[3]
        status = sys.argv[4] if len(sys.argv) > 4 else None
        log_entry(task, update, status)
    else:
        print(__doc__)
